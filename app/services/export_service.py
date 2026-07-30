from __future__ import annotations

import re
from pathlib import Path

from app.models.settings import FocusFieldConfig, TemplateConfig, TextStyleConfig, hao_to_pt
from app.services.bundled_fonts import normalize_font_family
from app.services.title_classifier import classify_paragraphs


INVALID_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|]+')
ENGLISH_INLINE_SPACING = re.compile(r"(?<=[A-Za-z0-9])([,;])(?=[A-Za-z])")
CHINESE_INLINE_SPACING = re.compile(r"(?<=[，。；：？！])\s+(?=[\u3400-\u9fff])")
CHINESE_SPACING_PUNCTUATION = set("，。；：？！、（）【】｛｝《》“”‘’")
CONTEXT_PUNCTUATION = set(
    ",，.。．;；:：?？！!()（）[]【】{}｛｝<>《》\"“”'‘’、"
)
CHINESE_PUNCTUATION = {
    ",": "，",
    "，": "，",
    ".": "。",
    "。": "。",
    "．": "。",
    ";": "；",
    "；": "；",
    ":": "：",
    "：": "：",
    "?": "？",
    "？": "？",
    "!": "！",
    "！": "！",
    "(": "（",
    "（": "（",
    ")": "）",
    "）": "）",
    "[": "【",
    "【": "【",
    "]": "】",
    "】": "】",
    "{": "｛",
    "｛": "｛",
    "}": "｝",
    "｝": "｝",
    "<": "《",
    "《": "《",
    ">": "》",
    "》": "》",
}
ENGLISH_PUNCTUATION = {
    "，": ",",
    "、": ",",
    "。": ".",
    "．": ".",
    "；": ";",
    "：": ":",
    "？": "?",
    "！": "!",
    "（": "(",
    "）": ")",
    "【": "[",
    "】": "]",
    "｛": "{",
    "｝": "}",
    "《": "<",
    "》": ">",
}


def export_docx_document(
    text: str,
    template: TemplateConfig,
    save_directory: str,
    focus_fields: FocusFieldConfig | None = None,
) -> str:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.shared import Mm, Pt
    except ImportError as exc:
        raise RuntimeError("缺少 python-docx 依赖，暂时无法导出 .docx 文件。") from exc

    normalized_text = normalize_punctuation_for_output(text)
    paragraphs = classify_paragraphs(normalized_text)
    if not paragraphs:
        raise ValueError("没有可导出的正文内容。")
    if template.clean_body_spaces:
        for paragraph_data in paragraphs:
            if paragraph_data["type"] == "body":
                paragraph_data["text"] = normalize_body_spacing_for_output(paragraph_data["text"])

    target_path = _build_export_path(paragraphs, save_directory)

    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(template.margins_mm.top)
    section.bottom_margin = Mm(template.margins_mm.bottom)
    section.left_margin = Mm(template.margins_mm.left)
    section.right_margin = Mm(template.margins_mm.right)

    style_map = {
        "title": template.title,
        "h1": template.h1,
        "h2": template.h2,
        "body": template.body,
    }

    body_indent = Pt(hao_to_pt(template.body.font_size_hao) * 2)
    h1_indent = Pt(hao_to_pt(template.h1.font_size_hao) * 2)
    h2_indent = Pt(hao_to_pt(template.h2.font_size_hao) * 2)

    for index, paragraph_data in enumerate(paragraphs):
        paragraph = document.add_paragraph()
        paragraph_type = paragraph_data["type"]
        config = style_map[paragraph_type]

        paragraph.paragraph_format.line_spacing = Pt(template.line_spacing_pt)
        paragraph.paragraph_format.first_line_indent = None
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)

        if paragraph_type == "title":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif paragraph_type == "body":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        else:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

        if paragraph_type == "body":
            paragraph.paragraph_format.first_line_indent = body_indent
        elif paragraph_type == "h1":
            paragraph.paragraph_format.first_line_indent = h1_indent
        elif paragraph_type == "h2":
            paragraph.paragraph_format.first_line_indent = h2_indent

        paragraph_text = paragraph_data["text"].lstrip() if paragraph_type == "title" else paragraph_data["text"]
        _add_text_runs(paragraph, paragraph_text, config, focus_fields, Pt, qn)

        if index == 0 and paragraph_type == "title" and len(paragraphs) > 1:
            blank_paragraph = document.add_paragraph()
            blank_paragraph.paragraph_format.line_spacing = Pt(template.line_spacing_pt)
            blank_paragraph.paragraph_format.first_line_indent = body_indent
            blank_paragraph.paragraph_format.space_before = Pt(0)
            blank_paragraph.paragraph_format.space_after = Pt(0)
            blank_paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    document.save(str(target_path))
    return str(target_path)


def _add_text_runs(paragraph, text: str, base_style: TextStyleConfig, focus_fields: FocusFieldConfig | None, pt_factory, qn) -> None:
    focus_values = _normalize_focus_values(focus_fields)
    if not focus_fields or not focus_values:
        run = paragraph.add_run(text)
        _apply_run_style(run, base_style, pt_factory, qn)
        return

    for segment_text, is_focus in _split_focus_segments(text, focus_values):
        run = paragraph.add_run(segment_text)
        if is_focus:
            _apply_run_style(run, focus_fields.style, pt_factory, qn, bold=focus_fields.bold)
        else:
            _apply_run_style(run, base_style, pt_factory, qn)


def _apply_run_style(run, style: TextStyleConfig, pt_factory, qn, *, bold: bool | None = None) -> None:
    font_family = normalize_font_family(style.font_family)
    run.font.name = font_family
    run.font.size = pt_factory(hao_to_pt(style.font_size_hao))
    if bold is not None:
        run.bold = bold

    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    r_fonts.set(qn("w:eastAsia"), font_family)


def _normalize_focus_values(focus_fields: FocusFieldConfig | None) -> list[str]:
    if focus_fields is None:
        return []

    values: list[str] = []
    for field in sorted(focus_fields.fields, key=len, reverse=True):
        normalized = field.strip()
        if normalized and normalized not in values:
            values.append(normalized)
    return values


def _split_focus_segments(text: str, focus_values: list[str]) -> list[tuple[str, bool]]:
    segments: list[tuple[str, bool]] = []
    index = 0
    while index < len(text):
        matched = next((field for field in focus_values if text.startswith(field, index)), None)
        if matched:
            segments.append((matched, True))
            index += len(matched)
            continue

        start = index
        while index < len(text) and not any(text.startswith(field, index) for field in focus_values):
            index += 1
        segments.append((text[start:index], False))

    return segments


def normalize_punctuation_for_output(text: str) -> str:
    normalized_lines = [_normalize_line_punctuation(line) for line in text.splitlines()]
    return "\n".join(normalized_lines)


def normalize_body_spacing_for_output(text: str) -> str:
    normalized_lines = [_normalize_body_line_spacing(line) for line in text.splitlines()]
    return "\n".join(normalized_lines)


def _normalize_body_line_spacing(text: str) -> str:
    normalized: list[str] = []
    index = 0

    while index < len(text):
        char = text[index]
        if not char.isspace():
            normalized.append(char)
            index += 1
            continue

        while index < len(text) and text[index].isspace():
            index += 1

        left = normalized[-1] if normalized else None
        right = text[index] if index < len(text) else None
        if left is None or right is None:
            continue
        if _should_remove_body_space(left, right):
            continue
        normalized.append(" ")

    return "".join(normalized)


def _should_remove_body_space(left: str, right: str) -> bool:
    if _is_chinese_char(left) and _is_chinese_char(right):
        return True
    return left in CHINESE_SPACING_PUNCTUATION or right in CHINESE_SPACING_PUNCTUATION


def _normalize_line_punctuation(text: str) -> str:
    normalized: list[str] = []
    double_quote_open = True
    single_quote_open = True

    for index, char in enumerate(text):
        if char in {'"', "“", "”"}:
            if _should_use_chinese_punctuation(text, index):
                normalized.append("“" if double_quote_open else "”")
                double_quote_open = not double_quote_open
            else:
                normalized.append('"')
            continue

        if char in {"'", "‘", "’"}:
            if _is_english_apostrophe(text, index):
                normalized.append("'")
            elif _should_use_chinese_punctuation(text, index):
                normalized.append("‘" if single_quote_open else "’")
                single_quote_open = not single_quote_open
            else:
                normalized.append("'")
            continue

        if char in CHINESE_PUNCTUATION or char in ENGLISH_PUNCTUATION or char == "、":
            normalized.append(_normalize_punctuation_char(text, index, char))
            continue

        normalized.append(char)

    normalized_text = ENGLISH_INLINE_SPACING.sub(r"\1 ", "".join(normalized))
    return CHINESE_INLINE_SPACING.sub("", normalized_text)


def _normalize_punctuation_char(text: str, index: int, char: str) -> str:
    if char in {".", "。", "．"} and _is_numeric_or_ordered_list_period(text, index):
        return "."
    if _should_use_chinese_punctuation(text, index):
        return CHINESE_PUNCTUATION.get(char, char)
    return ENGLISH_PUNCTUATION.get(char, char)


def _should_use_chinese_punctuation(text: str, index: int) -> bool:
    left = _nearest_context_char(text, index, -1)
    right = _nearest_context_char(text, index, 1)
    left_kind = _context_kind(left)
    right_kind = _context_kind(right)

    if left_kind == "english" and right_kind == "english":
        return False
    if right_kind == "english" and left_kind is None:
        return False
    if left_kind == "english" and right_kind is None:
        # 句尾英文产品名仍跟随整句语境，例如中文句中的 “Codepilot”。
        return _has_chinese_sentence_context(text, index)
    if left_kind == "chinese" or right_kind == "chinese":
        return True
    return False


def _nearest_context_char(text: str, index: int, step: int) -> str | None:
    cursor = index + step
    while 0 <= cursor < len(text):
        char = text[cursor]
        if char.isspace() or char in CONTEXT_PUNCTUATION:
            cursor += step
            continue
        return char
    return None


def _context_kind(char: str | None) -> str | None:
    if char is None:
        return None
    if _is_chinese_char(char):
        return "chinese"
    if char.isascii() and (char.isalpha() or char.isdigit()):
        return "english"
    return None


def _has_chinese_sentence_context(text: str, index: int) -> bool:
    cursor = index - 1
    while cursor >= 0:
        char = text[cursor]
        if char in "。.!?！？;；":
            return False
        if _is_chinese_char(char):
            return True
        cursor -= 1
    return False


def _is_chinese_char(char: str) -> bool:
    return (
        "\u3400" <= char <= "\u4dbf"
        or "\u4e00" <= char <= "\u9fff"
        or "\uf900" <= char <= "\ufaff"
    )


def _is_english_apostrophe(text: str, index: int) -> bool:
    left = text[index - 1] if index > 0 else ""
    right = text[index + 1] if index + 1 < len(text) else ""
    return left.isascii() and right.isascii() and left.isalpha() and right.isalpha()


def _is_numeric_or_ordered_list_period(text: str, index: int) -> bool:
    left = _nearest_context_char(text, index, -1)
    right = _nearest_context_char(text, index, 1)
    if left and right and left.isdigit() and right.isdigit():
        return True

    prefix = text[:index].strip()
    suffix = text[index + 1 :]
    return bool(prefix and len(prefix) <= 3 and prefix.isdigit() and suffix[:1].isspace())


def _build_export_path(paragraphs: list[dict[str, str]], save_directory: str) -> Path:
    directory = Path(save_directory).expanduser()
    if directory.suffix.lower() == ".docx":
        directory = directory.parent
    directory.mkdir(parents=True, exist_ok=True)

    original_title = paragraphs[0]["text"].strip() if paragraphs else ""
    safe_title = _sanitize_filename(original_title) or "排版结果"
    base_name = f"{safe_title}_排版结果"

    target_path = directory / f"{base_name}.docx"
    suffix_index = 1
    while target_path.exists():
        target_path = directory / f"{base_name}{suffix_index}.docx"
        suffix_index += 1

    return target_path


def _sanitize_filename(value: str) -> str:
    return INVALID_FILENAME_CHARS.sub("_", value).strip().strip(".")
