from __future__ import annotations

import tempfile
import unittest
from dataclasses import replace
from pathlib import Path

from docx import Document

from app.models.settings import (
    DEFAULT_TEMPLATE,
    FocusFieldConfig,
    TextStyleConfig,
    template_from_dict,
    template_to_dict,
)
from app.services.export_service import (
    export_docx_document,
    normalize_body_spacing_for_output,
    normalize_punctuation_for_output,
)


class BodySpacingNormalizationTests(unittest.TestCase):
    def test_removes_spaces_between_chinese_and_around_chinese_punctuation(self) -> None:
        source = "这 是　正\t文 ，\u00a0内  容。"

        self.assertEqual(
            "这是正文，内容。",
            normalize_body_spacing_for_output(source),
        )

    def test_collapses_other_whitespace_and_preserves_required_spaces(self) -> None:
        source = "使用  Microsoft\tWord  处理  100\u00a0kg"

        self.assertEqual(
            "使用 Microsoft Word 处理 100 kg",
            normalize_body_spacing_for_output(source),
        )

    def test_trims_line_edges_and_preserves_line_breaks(self) -> None:
        source = "  第一 行  \n\tEnglish   words\t"

        self.assertEqual(
            "第一行\nEnglish words",
            normalize_body_spacing_for_output(source),
        )


class TemplateConfigTests(unittest.TestCase):
    def test_legacy_template_defaults_to_enabled(self) -> None:
        template = template_from_dict({})

        self.assertTrue(template.clean_body_spaces)

    def test_template_round_trip_preserves_disabled_setting(self) -> None:
        template = replace(DEFAULT_TEMPLATE, clean_body_spaces=False)

        restored = template_from_dict(template_to_dict(template))

        self.assertFalse(restored.clean_body_spaces)

    def test_invalid_boolean_and_invalid_top_level_use_default(self) -> None:
        self.assertTrue(template_from_dict({"clean_body_spaces": "false"}).clean_body_spaces)
        self.assertTrue(template_from_dict([]).clean_body_spaces)


class ExportBodySpacingIntegrationTests(unittest.TestCase):
    def _export_and_read(
        self,
        text: str,
        *,
        clean_body_spaces: bool,
        focus_fields: FocusFieldConfig | None = None,
    ) -> Document:
        template = replace(DEFAULT_TEMPLATE, clean_body_spaces=clean_body_spaces)
        temporary_directory = tempfile.TemporaryDirectory()
        self.addCleanup(temporary_directory.cleanup)

        exported_path = export_docx_document(
            text,
            template,
            temporary_directory.name,
            focus_fields,
        )
        self.assertTrue(Path(exported_path).exists())
        return Document(exported_path)

    def test_export_cleans_only_body_paragraphs(self) -> None:
        document = self._export_and_read(
            "\n".join(
                (
                    "文 档 标 题",
                    "一、  一级 标题",
                    "（一）  二级 标题",
                    "这 是  Microsoft  Word  正文 ， 数量  100 kg。",
                )
            ),
            clean_body_spaces=True,
        )

        self.assertEqual(
            [
                "文 档 标 题",
                "",
                "一、  一级 标题",
                "（一）  二级 标题",
                "这是 Microsoft Word 正文，数量 100 kg。",
            ],
            [paragraph.text for paragraph in document.paragraphs],
        )

    def test_disabled_setting_preserves_body_spacing(self) -> None:
        document = self._export_and_read(
            "文档标题\n正  文",
            clean_body_spaces=False,
        )

        self.assertEqual("正  文", document.paragraphs[-1].text)

    def test_focus_field_matching_runs_after_spacing_cleanup(self) -> None:
        focus_fields = FocusFieldConfig(
            fields=["分析认为"],
            style=TextStyleConfig(font_family="仿宋", font_size_hao="小二"),
            bold=True,
        )
        document = self._export_and_read(
            "文档标题\n分 析 认 为，工作进展顺利。",
            clean_body_spaces=True,
            focus_fields=focus_fields,
        )
        body_paragraph = document.paragraphs[-1]

        self.assertEqual("分析认为，工作进展顺利。", body_paragraph.text)
        self.assertEqual("分析认为", body_paragraph.runs[0].text)
        self.assertTrue(body_paragraph.runs[0].bold)


class PunctuationRegressionTests(unittest.TestCase):
    def test_chinese_and_english_sentence_punctuation_still_adapts(self) -> None:
        self.assertEqual("中文句子，", normalize_punctuation_for_output("中文句子,"))
        self.assertEqual("English sentence.", normalize_punctuation_for_output("English sentence。"))

    def test_decimal_point_is_preserved(self) -> None:
        self.assertEqual("版本号 3.12", normalize_punctuation_for_output("版本号 3.12"))


if __name__ == "__main__":
    unittest.main()
